"""Tests for src/report.py, using tests/fixtures/expenses_broken.csv."""

from pathlib import Path

import pytest
from openpyxl import load_workbook

import report
import validate

FIXTURE = Path(__file__).parent / "fixtures" / "expenses_broken.csv"

SHEET_NAMES = [
    "Summary",
    "By Restaurant Type",
    "By Cuisine",
    "By Year",
    "Receipts",
    "Needs Review",
]


@pytest.fixture(scope="module")
def workbook(tmp_path_factory):
    """Write the report for the broken fixture once and open it."""
    clean_rows, issues = validate.validate_file(FIXTURE)
    path = tmp_path_factory.mktemp("report") / "food_report.xlsx"
    report.write_report(clean_rows, issues, path)
    return load_workbook(path)


def summary_values(workbook):
    """Return the Summary sheet as a {metric: value} dict."""
    values = {}
    for metric, value in workbook["Summary"].iter_rows(min_row=2, values_only=True):
        values[metric] = value
    return values


def sheet_rows(sheet):
    """Return the data rows of a sheet (no header) as tuples."""
    return list(sheet.iter_rows(min_row=2, values_only=True))


def test_sheet_names(workbook):
    assert workbook.sheetnames == SHEET_NAMES


def test_summary_numbers(workbook):
    values = summary_values(workbook)
    assert values["Number of receipts"] == 7
    assert values["Total spent"] == pytest.approx(259.38)
    assert values["Average per visit"] == pytest.approx(259.38 / 7)
    assert values["Receipts with a tip"] == 3
    # T003 8/40, T005 3/18, T006 7.38/41
    expected_tip = (8 / 40 + 3 / 18 + 7.38 / 41) / 3
    assert values["Average tip percentage (receipts with a tip)"] == pytest.approx(expected_tip)
    assert values["Receipts without a date (left out of By Year)"] == 1
    assert values["Rows excluded (errors)"] == 4
    assert values["Rows flagged for review (kept, with warnings)"] == 6


def test_by_restaurant_type(workbook):
    rows = sheet_rows(workbook["By Restaurant Type"])
    by_type = {}
    for name, visits, total, average in rows:
        by_type[name] = (visits, total, average)
    assert by_type["Casual Dining"][0] == 3
    assert by_type["Casual Dining"][1] == pytest.approx(56.58 + 69.25 + 51.60)
    assert by_type["Cafe"][0] == 2
    assert by_type["Other"][0] == 1
    assert by_type["Fast Food"][2] == pytest.approx(7.61)


def test_by_cuisine_sorted_by_total(workbook):
    rows = sheet_rows(workbook["By Cuisine"])
    totals = []
    for row in rows:
        totals.append(row[2])
    assert totals == sorted(totals, reverse=True)
    assert rows[0][0] == "American"


def test_by_year_leaves_out_empty_date(workbook):
    rows = sheet_rows(workbook["By Year"])
    assert rows[0] == (2015, 3, pytest.approx(69.25 + 51.60 + 39.18))
    assert rows[1] == (2016, 2, pytest.approx(56.58 + 7.61))
    assert rows[2] == (2017, 1, pytest.approx(22.62))
    assert len(rows) == 3


def test_receipts_sheet_has_every_clean_row(workbook):
    sheet = workbook["Receipts"]
    assert len(sheet_rows(sheet)) == 7
    header = []
    for cell in sheet[1]:
        header.append(cell.value)
    assert header == validate.COLUMNS + ["row", "year"]


def test_needs_review_has_every_issue(workbook):
    _, issues = validate.validate_file(FIXTURE)
    sheet = workbook["Needs Review"]
    assert len(sheet_rows(sheet)) == len(issues)
    header = []
    for cell in sheet[1]:
        header.append(cell.value)
    assert header == validate.ISSUE_FIELDS


def test_headers_are_bold(workbook):
    for name in SHEET_NAMES:
        for cell in workbook[name][1]:
            assert cell.font.bold, f"{name} header {cell.value} is not bold"


def test_money_is_currency_formatted(workbook):
    assert workbook["Summary"]["B3"].number_format == report.MONEY_FORMAT
    assert workbook["By Year"]["C2"].number_format == report.MONEY_FORMAT
    receipts = workbook["Receipts"]
    total_column = validate.COLUMNS.index("total") + 1
    assert receipts.cell(row=2, column=total_column).number_format == report.MONEY_FORMAT


def test_column_widths_are_set(workbook):
    sheet = workbook["Summary"]
    assert sheet.column_dimensions["A"].width > 30


def make_row(receipt_id, subtotal, tip, total):
    """Return a valid row dict with the given money values and no tax."""
    return {
        "receipt_id": receipt_id,
        "date": "2026-10-02",
        "merchant": "Chipotle",
        "restaurant_type": "Fast Food",
        "cuisine": "Mexican",
        "description": "Burrito bowl",
        "subtotal": subtotal,
        "tax": "0.00",
        "tip": tip,
        "total": total,
        "currency": "USD",
        "payment_method": "Card",
        "confidence": "high",
        "notes": "",
    }


def test_zero_tip_counts_and_empty_tip_is_left_out():
    rows = [
        make_row("X001", "20.00", "0.00", "20.00"),  # tip line showing zero
        make_row("X002", "20.00", "5.00", "25.00"),  # 25% tip
        make_row("X003", "10.00", "", "10.00"),  # no tip line
    ]
    clean_rows, issues = validate.validate_rows(rows)
    assert issues == []

    tips = report.tip_percentages(clean_rows)
    assert list(tips) == [0.0, 0.25]

    summary = {}
    for metric, value, _ in report.build_summary(clean_rows, issues):
        summary[metric] = value
    assert summary["Receipts with a tip"] == 2
    assert summary["Average tip percentage (receipts with a tip)"] == pytest.approx(0.125)


def test_empty_data_does_not_crash(tmp_path):
    csv_path = tmp_path / "header_only.csv"
    csv_path.write_text(",".join(validate.COLUMNS) + "\n")
    clean_rows, issues = validate.validate_file(csv_path)
    path = report.write_report(clean_rows, issues, tmp_path / "empty.xlsx")
    values = summary_values(load_workbook(path))
    assert values["Number of receipts"] == 0
    assert values["Average per visit"] is None

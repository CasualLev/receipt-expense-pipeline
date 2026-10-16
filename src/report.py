"""Build the Excel spending report from the validated expense rows.

Use write_report(clean_rows, issues) with the output of validate.validate_file.
"""

import sys
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

import validate

DEFAULT_OUTPUT = Path("output") / "food_report.xlsx"

MONEY_FORMAT = '"$"#,##0.00'
PERCENT_FORMAT = "0.0%"
MONEY_HEADERS = ["subtotal", "tax", "tip", "total", "Total spent", "Average per visit"]
MAX_COLUMN_WIDTH = 60


# ---------------------------------------------------------------------------
# Calculations
# ---------------------------------------------------------------------------


def tip_percentages(clean_rows):
    """Return tip / subtotal for each receipt with a tip above zero."""
    has_tip = clean_rows["tip"] > 0
    has_subtotal = clean_rows["subtotal"] > 0
    with_tip = clean_rows[has_tip & has_subtotal]
    return with_tip["tip"] / with_tip["subtotal"]


def flagged_row_count(clean_rows, issues):
    """Count kept rows that have at least one warning."""
    kept_rows = set(clean_rows["row"])
    flagged = set()
    for issue in issues:
        if issue.severity == validate.WARNING and issue.row in kept_rows:
            flagged.add(issue.row)
    return len(flagged)


def average_or_none(values):
    """Return the mean of values, or None when there are none."""
    if len(values) == 0:
        return None
    return float(values.mean())


def build_summary(clean_rows, issues):
    """Return the Summary sheet as rows of (metric, value, number format)."""
    tips = tip_percentages(clean_rows)
    no_date = clean_rows["year"].isna().sum()
    return [
        ("Number of receipts", len(clean_rows), None),
        ("Total spent", float(clean_rows["total"].sum()), MONEY_FORMAT),
        ("Average per visit", average_or_none(clean_rows["total"]), MONEY_FORMAT),
        ("Average tip percentage (receipts with a tip)", average_or_none(tips), PERCENT_FORMAT),
        ("Receipts with a tip", len(tips), None),
        ("Receipts without a date (left out of By Year)", int(no_date), None),
        ("Rows excluded (errors)", len(validate.excluded_rows(issues)), None),
        ("Rows flagged for review (kept, with warnings)", flagged_row_count(clean_rows, issues), None),
    ]


def build_breakdown(clean_rows, column, label):
    """Return visits, total spent and average per visit for each value of column."""
    grouped = clean_rows.groupby(column)["total"]
    table = pd.DataFrame(
        {
            "Visits": grouped.size(),
            "Total spent": grouped.sum(),
            "Average per visit": grouped.mean(),
        }
    )
    table = table.sort_values("Total spent", ascending=False)
    table.index.name = label
    return table.reset_index()


def build_by_year(clean_rows):
    """Return visits and total spent per year. Rows with no date are skipped."""
    grouped = clean_rows.groupby("year")["total"]
    table = pd.DataFrame({"Visits": grouped.size(), "Total spent": grouped.sum()})
    table.index.name = "Year"
    return table.reset_index()


# ---------------------------------------------------------------------------
# Formatting
# ---------------------------------------------------------------------------


def format_sheet(sheet):
    """Bold the header row, format money columns and set column widths."""
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    for column_cells in sheet.iter_cols(min_row=1):
        header = column_cells[0].value
        if header in MONEY_HEADERS:
            for cell in column_cells[1:]:
                cell.number_format = MONEY_FORMAT

    set_column_widths(sheet)


def display_length(cell):
    """Roughly how many characters a cell shows in Excel."""
    if cell.value is None:
        return 0
    if cell.number_format == MONEY_FORMAT and isinstance(cell.value, (int, float)):
        return len(f"${cell.value:,.2f}")
    if cell.number_format == PERCENT_FORMAT:
        return 6
    return len(str(cell.value))


def set_column_widths(sheet):
    """Make each column wide enough for its longest value, up to a limit."""
    for index, column_cells in enumerate(sheet.iter_cols(min_row=1), start=1):
        longest = 0
        for cell in column_cells:
            longest = max(longest, display_length(cell))
        width = min(longest + 3, MAX_COLUMN_WIDTH)
        sheet.column_dimensions[get_column_letter(index)].width = width


def write_summary_sheet(writer, summary_rows):
    """Write the Summary sheet and give each value its own number format."""
    pairs = []
    for metric, value, _ in summary_rows:
        pairs.append((metric, value))
    table = pd.DataFrame(pairs, columns=["Metric", "Value"])
    table.to_excel(writer, sheet_name="Summary", index=False)
    sheet = writer.sheets["Summary"]

    for index, (_, _, number_format) in enumerate(summary_rows):
        if number_format is not None:
            sheet.cell(row=index + 2, column=2).number_format = number_format
    format_sheet(sheet)


def write_table_sheet(writer, table, sheet_name):
    """Write one DataFrame to its own formatted sheet."""
    table.to_excel(writer, sheet_name=sheet_name, index=False)
    format_sheet(writer.sheets[sheet_name])


# ---------------------------------------------------------------------------
# Main entry points
# ---------------------------------------------------------------------------


def write_report(clean_rows, issues, output_path=DEFAULT_OUTPUT):
    """Create the Excel report and return its path."""
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        write_summary_sheet(writer, build_summary(clean_rows, issues))
        write_table_sheet(
            writer,
            build_breakdown(clean_rows, "restaurant_type", "Restaurant type"),
            "By Restaurant Type",
        )
        write_table_sheet(writer, build_breakdown(clean_rows, "cuisine", "Cuisine"), "By Cuisine")
        write_table_sheet(writer, build_by_year(clean_rows), "By Year")
        write_table_sheet(writer, clean_rows, "Receipts")
        write_table_sheet(writer, validate.issues_to_dataframe(issues), "Needs Review")

    return output_path


def main():
    """Validate a CSV and write the report: python src/report.py data/expenses.csv"""
    if len(sys.argv) != 2:
        print("Usage: python src/report.py path/to/expenses.csv")
        return 1

    clean_rows, issues = validate.validate_file(sys.argv[1])
    try:
        path = write_report(clean_rows, issues)
    except PermissionError:
        print(f"Could not write {DEFAULT_OUTPUT}. Is it open in Excel?")
        return 1
    print(f"Report written to {path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
